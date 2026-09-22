"""One-time extraction: URPS_Didactic_Schedule.xlsx + GTL mapping PDF -> data/curriculum.json.

After this runs, data/curriculum.json is the canonical source; edit it, not the xlsx/PDF.
"""

import json
import re
import sys
import urllib.request
from pathlib import Path

import openpyxl
import pymupdf

SRC = Path("/mnt/c/Users/berkl/OneDrive/Fellowship/Education/Schedule")
OUT = Path(__file__).resolve().parent.parent / "data" / "curriculum.json"

GTL_SECTIONS = [
    ("anatomy", "Anatomy & Physiology"),
    ("ui-luts", "UI & LUTS"),
    ("pop", "Pelvic Organ Prolapse"),
    ("fi-dd", "FI & Defecatory Dysfunction"),
    ("rvf", "Rectovaginal Fistulas"),
    ("ugf-ud", "UG Fistulas & Urethral Diverticula"),
    ("gsm", "GSM & Vulvar Conditions"),
    ("uti", "UTI & Hematuria"),
    ("neuro", "Neuro-Urology"),
    ("pbs", "Pelvic Pain & PBS"),
    ("periop", "Peri-Operative Management"),
    ("scholarly", "Scholarly Activity & QI"),
    ("enrichment", "Professional & Enrichment"),
]
CLINICAL_GTL = [k for k, _ in GTL_SECTIONS if k not in ("scholarly", "enrichment")]

GTL_ALIASES = {
    "pelvic floor anatomy & physiology": ["anatomy"],
    "ui & luts": ["ui-luts"],
    "pop treatment": ["pop"],
    "pop": ["pop"],
    "fi & dd treatment": ["fi-dd"],
    "fi & dd": ["fi-dd"],
    "rectovaginal fistulas": ["rvf"],
    "rvf": ["rvf"],
    "ug fistulas & urethral diverticula": ["ugf-ud"],
    "ug fistulas & ud": ["ugf-ud"],
    "ug fistulas": ["ugf-ud"],
    "gsm & other vulvar conditions": ["gsm"],
    "uti & hematuria": ["uti"],
    "uti": ["uti"],
    "neuro-urology & neurogenic lutd": ["neuro"],
    "pbs & pelvic floor dysfunction": ["pbs"],
    "pbs": ["pbs"],
    "general peri-operative management": ["periop"],
    "general peri-op mgmt": ["periop"],
    "peri-op mgmt": ["periop"],
    "scholarly activity": ["scholarly"],
    "qi science": ["scholarly"],
    "scholarly activity — qi science": ["scholarly"],
    "scholarly activity — dissemination": ["scholarly"],
    "program enrichment (gtl foreword)": ["enrichment"],
    "acgme common program requirements": ["enrichment"],
    "cross-content (all clinical sections)": CLINICAL_GTL,
    "all gtl sections": [k for k, _ in GTL_SECTIONS if k != "enrichment"],
    "fi & dd (shared framework)": ["fi-dd"],
}

PLAIN_HEADINGS = {"Program Enrichment", "Post-Cancer PFD", "Comprehensive Review",
                  "Professional Development", "ACGME Competencies"}

TEXT_FIXES = {"FI■": "FI&DD", "FII ": "FI&DD ", "Kurtulu■": "Kurtuluş", "UF&UD;)": "UF&UD)"}


def fix(s: str) -> str:
    for a, b in TEXT_FIXES.items():
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", s).strip()


def gtl_keys(label: str) -> list[str]:
    keys: list[str] = []
    for part in label.split(";"):
        p = part.strip().lower()
        if p not in GTL_ALIASES:
            sys.exit(f"unknown GTL section: {part!r}")
        keys += [k for k in GTL_ALIASES[p] if k not in keys]
    return keys


def parse_pdf():
    doc = pymupdf.open(SRC / "URPS_GTL_Learning_Objectives_Mapping.pdf")
    lines = []
    for page in doc:
        for ln in page.get_text().splitlines():
            ln = ln.strip()
            if not ln or ln.startswith("URPS Fellowship Didactic Curriculum") or re.fullmatch(r"Page \d+", ln):
                continue
            lines.append(ln)

    blocks, sessions = {}, {}
    block_no, i = None, 0
    while i < len(lines):
        m = re.fullmatch(r"0(\d) YEAR (\d) — BLOCK \d: (.+)", lines[i])
        if m:
            block_no = int(m.group(1))
            blocks[block_no] = {"number": block_no, "year": int(m.group(2)), "name": re.sub(r"\bIi\b", "II", m.group(3).title())}
            i += 1
            continue
        m = re.fullmatch(r"SESSION (\d+)", lines[i])
        if not m:
            i += 1
            continue
        n = int(m.group(1))
        title = lines[i + 2]
        i += 3
        gtl_label = None
        if lines[i].startswith("GTL Section:"):
            gtl_label = fix(lines[i].removeprefix("GTL Section:"))
            i += 1
        groups, desc, cur, bullet = [], [], None, None
        while lines[i] != "READING":
            ln = lines[i]
            if ln == "–":
                if cur is None:
                    cur = {"heading": None, "objectives": []}
                    groups.append(cur)
                bullet = []
                cur["objectives"].append(bullet)
            elif re.search(r"\(GTL[^)]*\)$", ln) or ln in PLAIN_HEADINGS:
                cur = {"heading": fix(ln), "objectives": []}
                groups.append(cur)
                bullet = None
            elif bullet is not None:
                bullet.append(ln)
            else:
                desc.append(ln)
            i += 1
        for g in groups:
            g["objectives"] = [fix(" ".join(b)) for b in g["objectives"]]
        sessions[n] = {"block": block_no, "title": title, "gtlLabel": gtl_label,
                       "groups": groups, "description": fix(" ".join(desc)) or None}
        i += 1
    return blocks, sessions


def split_citations(cell):
    if not cell:
        return []
    parts = re.split(r"(?<=\))\s*;\s*", str(cell).strip())
    out = []
    for p in parts:
        p = p.strip()
        pmid = re.search(r"PMID:\s*(\d+)", p)
        label = re.match(r"^([A-Z][A-Za-z0-9-]+(?: [A-Z][A-Za-z0-9-]+)? (?:Trial|Study)):\s*", p)
        out.append({
            "citation": p,
            "label": label.group(1) if label else None,
            "pmid": pmid.group(1) if pmid else None,
        })
    return out


def walters_chapters(reading: str) -> list[int]:
    m = re.search(r"Walters Ch\. ([\d,\s-]+)", reading)
    if not m:
        return []
    chs = []
    for part in m.group(1).split(","):
        part = part.strip().rstrip("-")
        if not part:
            continue
        if "-" in part:
            a, b = map(int, part.split("-"))
            chs += list(range(a, b + 1))
        else:
            chs.append(int(part))
    return chs


def pubmed_meta(pmids: list[str]) -> dict:
    url = ("https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi?db=pubmed&retmode=json&id="
           + ",".join(pmids))
    with urllib.request.urlopen(url, timeout=60) as r:
        res = json.load(r)["result"]
    meta = {}
    for pmid in pmids:
        rec = res.get(pmid)
        if not rec or "error" in rec:
            print(f"warning: PubMed has no record for {pmid}", file=sys.stderr)
            continue
        pmc = next((a["value"] for a in rec.get("articleids", []) if a["idtype"] == "pmc"), None)
        meta[pmid] = {
            "title": rec.get("title", "").rstrip("."),
            "journal": rec.get("source"),
            "year": int(rec["pubdate"][:4]) if rec.get("pubdate", "")[:4].isdigit() else None,
            "firstAuthor": (rec.get("authors") or [{}])[0].get("name"),
            "pmcid": pmc,
        }
    return meta


def main():
    blocks, pdf = parse_pdf()
    wb = openpyxl.load_workbook(SRC / "URPS_Didactic_Schedule.xlsx", data_only=True)

    urp = {}
    for row in wb["URP Code Index"].iter_rows(min_row=4, values_only=True):
        if row[0] and re.fullmatch(r"URP\d", str(row[0])):
            urp[row[0]] = {"code": row[0], "name": row[1], "description": row[2]}

    sessions = []
    for row in wb["Complete Schedule"].iter_rows(min_row=2, values_only=True):
        if not isinstance(row[0], int):
            continue
        n, _year, _block, codes, topic, reading, primary, further = row[:8]
        p = pdf[n]
        if p["title"].strip() != topic.strip():
            sys.exit(f"session {n}: xlsx topic {topic!r} != PDF {p['title']!r}")
        codes = None if codes == "N/A" else codes
        kind = "mock" if topic.startswith("Mock Oral Boards") else ("clinical" if codes else "professional")  # kind is now hand-maintained in data/curriculum.json (sessions 19, 22, 33, 36, 37, 40 are clinical)
        reading = re.sub(r"Review all Year 1 Spring topics", "Review all Block 1 topics", reading)
        reading = re.sub(r"Review all Year 1 Fall topics", "Review all Block 2 topics", reading)
        sessions.append({
            "number": n,
            "block": p["block"],
            "kind": kind,
            "title": topic,
            "urpCodes": codes,
            "urpDomains": sorted(set(re.findall(r"URP\d", codes or ""))),
            "gtlLabel": p["gtlLabel"],
            "gtl": gtl_keys(p["gtlLabel"]) if p["gtlLabel"] else [],
            "description": p["description"],
            "objectiveGroups": p["groups"],
            "reading": reading,
            "waltersChapters": walters_chapters(reading),
            "primary": split_citations(primary),
            "further": split_citations(further),
        })

    # Mock boards cover every GTL section taught in their block (the final one covers all).
    for s in sessions:
        if s["kind"] == "mock":
            pool = sessions if s["number"] == 59 else [x for x in sessions if x["block"] == s["block"]]
            s["gtl"] = [k for k, _ in GTL_SECTIONS if any(k in x["gtl"] for x in pool if x is not s)]

    pmids = sorted({a["pmid"] for s in sessions for a in s["primary"] + s["further"] if a["pmid"]}, key=int)
    meta = pubmed_meta(pmids)
    for s in sessions:
        for a in s["primary"] + s["further"]:
            if a["pmid"] in meta:
                a["pubmed"] = meta[a["pmid"]]

    for b in blocks.values():
        ss = [s for s in sessions if s["block"] == b["number"]]
        b["sessions"] = [ss[0]["number"], ss[-1]["number"]]

    data = {
        "title": "URPS Fellowship Curriculum",
        "blocks": list(blocks.values()),
        "gtlSections": [{"key": k, "name": v} for k, v in GTL_SECTIONS],
        "urpDomains": list(urp.values()),
        "sessions": sessions,
    }
    OUT.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n")
    kinds = {k: sum(s["kind"] == k for s in sessions) for k in ("clinical", "professional", "mock")}
    print(f"wrote {OUT}: {len(sessions)} sessions {kinds}, {len(pmids)} PMIDs, {len(meta)} with metadata")


if __name__ == "__main__":
    main()
